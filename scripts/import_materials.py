#!/usr/bin/env python3
"""Builds your local materials library from Materials_Hunt_List.xlsx.

    python scripts/import_materials.py path/to/Materials_Hunt_List.xlsx [--output FILE]

Writes your user library (puzzle.materials.user_library_path(), e.g.
~/Library/Application Support/Puzzle/materials/library.json on macOS) unless
--output names another file; Puzzle loads it on top of its generic materials. The
library stays on your machine; it is not part of the Puzzle repository. The
workbook is shared with the Tension_Meter app; edit the workbook and re-run, don't
edit the JSON by hand.
It must have been recalculated (opened and saved), because derived SI columns
are formulas and this reads their cached values. Needs openpyxl.

Read:
- Materials tab: room-temperature modulus, density, yield and ultimate strength
  (typed imperial columns, converted exactly), type, group, specification,
  basis and notes.
- Temperature Data tab: temperature-dependent values in SI, each with its source.
  "% of RT" rows become absolute values using the room-temperature figure.
  Mean CTE rows keep their reference temperature; total thermal expansion from
  293 K becomes a mean CTE from 293.15 K when no mean CTE is given.
- Sources tab: links for each source title.

Not read yet: fatigue, stress-strain models, weld fillers, material links.
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOM_K = 293.15

PA_PER_MSI = 6894757.293168e3
PA_PER_KSI = 6894757.293168
KG_M3_PER_LB_IN3 = 0.45359237 / 0.0254**3

# Workbook property -> (Puzzle table key, SI unit, factor from the workbook's SI unit).
PROPERTIES = {
    "Modulus of elasticity": ("youngs_modulus", "Pa", 1e9),
    "Shear modulus": ("shear_modulus", "Pa", 1e9),
    "Poisson's ratio": ("poissons_ratio", "dimensionless", 1.0),
    "Yield strength": ("yield_strength", "Pa", 1e6),
    "Tensile strength": ("ultimate_strength", "Pa", 1e6),
    "Elongation": ("elongation", "percent", 1.0),
    "Thermal conductivity": ("thermal_conductivity", "W/(m*K)", 1.0),
    "Specific heat": ("specific_heat", "J/(kg*K)", 1.0),
    "Thermal diffusivity": ("thermal_diffusivity", "m^2/s", 1e-6),
    "Electrical resistivity": ("electrical_resistivity", "ohm*m", 1e-8),
}
# "% of RT" rows scale the Materials-tab value of the same property.
PERCENT_OF_RT = {
    "Modulus of elasticity, % of RT": "youngs_modulus",
    "Yield strength, % of RT": "yield_strength",
    "Tensile strength, % of RT": "ultimate_strength",
}
# Mean CTE reference temperatures, most preferred first.
MEAN_CTE = {
    "Mean CTE from 20 °C": 293.15,
    "Mean CTE from 70 °F": 294.261,
    "Mean CTE from 0 °C": 273.15,
}
TOTAL_EXPANSION = "Thermal expansion from 293 K"
UNITS = {key: unit for key, unit, _ in PROPERTIES.values()} | {"thermal_expansion": "1/K"}


def rows(ws):
    header = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    for values in ws.iter_rows(min_row=5, values_only=True):
        row = {
            h: (v.strip() if isinstance(v, str) else v)
            for h, v in zip(header, values, strict=False)
            if h
        }
        if row.get("Descriptor"):
            yield row


def slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def number(value):
    if value in (None, ""):
        return None
    if isinstance(value, str) and value.startswith("="):
        raise SystemExit(f"formula where a value was expected: {value}")
    return float(value)


def main(xlsx: Path, out: Path) -> None:
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    links = {
        r[1].strip(): (r[2] or "").strip()
        for r in wb["Sources"].iter_rows(min_row=5, values_only=True)
        if r[1]
    }

    materials: dict[str, dict] = {}
    for r in rows(wb["Materials"]):
        e, rho = number(r["Mod of E. (msi)"]), number(r["density(lb/in^3)"])
        fy, fu = number(r["Yield Strength (ksi)"]), number(r["Ultimate Strength (ksi)"])
        props = {
            "density": [rho * KG_M3_PER_LB_IN3, "kg/m^3"],
            "youngs_modulus": [e * PA_PER_MSI, "Pa"],
        }
        if fy is not None:
            props["yield_strength"] = [fy * PA_PER_KSI, "Pa"]
        if fu is not None:
            props["ultimate_strength"] = [fu * PA_PER_KSI, "Pa"]
        info = {
            k: r.get(c) or ""
            for k, c in (
                ("type", "Material Type"),
                ("group", "Material Group"),
                ("specification", "Specification"),
                ("basis", "Basis"),
            )
        }
        materials[r["Descriptor"]] = {
            "name": r["Descriptor"],
            "description": r.get("Notes") or "",
            "info": {k: v for k, v in info.items() if v},
            "properties": props,
            "points": defaultdict(list),
            "cte": defaultdict(list),
            "total": [],
        }

    skipped = defaultdict(int)
    for r in rows(wb["Temperature Data"]):
        m = materials.get(r["Descriptor"])
        if m is None:
            skipped[r["Descriptor"]] += 1
            continue
        prop, k, v = (
            r["Property"],
            number(r["Temperature (K)"]),
            number(r["Value (SI)"]),
        )
        if k is None or v is None:
            raise SystemExit(
                "Temperature Data has no cached values - recalculate the workbook first"
            )
        src = (r.get("Source") or "").strip()
        if prop in PROPERTIES:
            key, _, factor = PROPERTIES[prop]
            m["points"][key].append((k, v * factor, src))
        elif prop in PERCENT_OF_RT:
            key = PERCENT_OF_RT[prop]
            base = m["properties"].get(key)
            if base:
                m["points"][key].append((k, v / 100 * base[0], src))
        elif prop in MEAN_CTE:
            m["cte"][prop].append((k, v * 1e-6, src))
        elif prop == TOTAL_EXPANSION:
            m["total"].append((k, v / 100, src))
        else:
            raise SystemExit(f"unknown property in Temperature Data: {prop}")

    library = {}
    for name, m in materials.items():
        sources: list[str] = []

        def source_index(title: str, sources: list[str] = sources):
            if not title:
                return None
            label = f"{title} ({links[title]})" if links.get(title) else title
            if label not in sources:
                sources.append(label)
            return sources.index(label)

        tables = {}
        for key, pts in m["points"].items():
            tables[key] = {
                "unit": UNITS[key],
                "points": [[k, v, source_index(s)] for k, v, s in sorted(pts)],
            }
        cte = (
            next((MEAN_CTE[p], m["cte"][p]) for p in MEAN_CTE if m["cte"].get(p))
            if m["cte"]
            else None
        )
        if cte is None and m["total"]:
            pts = [
                (k, strain / (k - ROOM_K), s) for k, strain, s in m["total"] if abs(k - ROOM_K) > 5
            ]
            cte = (ROOM_K, pts) if pts else None
        if cte:
            ref, pts = cte
            tables["thermal_expansion"] = {
                "unit": "1/K",
                "reference_temperature": ref,
                "points": [[k, v, source_index(s)] for k, v, s in sorted(pts)],
            }

        if "specific_heat" not in tables and {"thermal_conductivity", "thermal_diffusivity"} <= set(
            tables
        ):
            derived = derive_specific_heat(tables, m["properties"]["density"][0])
            if derived:
                label = source_index(DERIVED_CP)
                tables["specific_heat"] = {
                    "unit": "J/(kg*K)",
                    "points": [[k, v, label] for k, v in derived],
                }

        props = dict(m["properties"])
        for key, table in tables.items():
            if key == "thermal_expansion" and key not in props:
                # Datasheets quote mean CTE over a range from the reference (e.g. 20-100 °C);
                # the narrowest range on offer is the room-temperature value.
                ref = table["reference_temperature"]
                props[key] = [min(table["points"], key=lambda pt: abs(pt[0] - ref))[1], "1/K"]
            elif key not in props and key not in ("elongation", "thermal_diffusivity"):
                rt = room_value(table["points"])
                if rt is not None:
                    props[key] = [rt, table["unit"]]
        if "electrical_resistivity" in props:
            rho_e = props.pop("electrical_resistivity")[0]
            props["electrical_conductivity"] = [1 / rho_e, "S/m"]

        library[slug(name)] = {
            "name": name,
            "description": m["description"],
            "info": m["info"],
            "properties": {
                k: {"value": round_sig(v), "unit": u} for k, (v, u) in sorted(props.items())
            },
            "tables": {
                k: {**t, "points": [[round_sig(a), round_sig(b), s] for a, b, s in t["points"]]}
                for k, t in sorted(tables.items())
            },
            "sources": sources,
        }

    if len(library) != len(materials):
        raise SystemExit("two materials share a key; rename one in the workbook")
    write(library, xlsx, out)
    n_points = sum(len(t["points"]) for m in library.values() for t in m["tables"].values())
    print(f"wrote {out}: {len(library)} materials, {n_points} table points")
    for name, n in sorted(skipped.items()):
        print(f"  skipped {n} temperature rows for '{name}' (not on the Materials tab)")


DERIVED_CP = (
    "Derived: specific heat = thermal conductivity / (density x thermal diffusivity), "
    "from this material's conductivity and diffusivity sources, at room-temperature density"
)


def derive_specific_heat(tables, density):
    """cp(T) = k(T) / (rho * alpha(T)) at each diffusivity temperature inside the k table."""
    k_curve = sorted((k, v) for k, v, _ in tables["thermal_conductivity"]["points"])
    out = []
    for t, alpha, _ in sorted(tables["thermal_diffusivity"]["points"]):
        k = interpolate(k_curve, t)
        if k is not None and alpha > 0:
            out.append((t, k / (density * alpha)))
    return out


def interpolate(curve, t):
    for (k0, v0), (k1, v1) in zip(curve, curve[1:], strict=False):
        if k0 <= t <= k1 and k1 > k0:
            return v0 + (v1 - v0) * (t - k0) / (k1 - k0)
    return next((v for k, v in curve if k == t), None)


def room_value(points):
    """The value at 293.15 K: interpolated if the table spans it, else a point within 10 K."""
    curve = sorted((k, v) for k, v, _ in points)
    inside = interpolate(curve, ROOM_K)
    if inside is not None:
        return inside
    near = [v for k, v in curve if abs(k - ROOM_K) <= 10]
    return sum(near) / len(near) if near else None


def round_sig(x, sig=7):
    return float(f"{x:.{sig}g}") if isinstance(x, float) else x


def write(library, xlsx, out: Path):
    lines = ["{", f' "generated_from": {json.dumps(xlsx.name)},', ' "materials": {']
    items = list(library.items())
    for i, (key, m) in enumerate(items):
        comma = "," if i < len(items) - 1 else ""
        lines.append(f"  {json.dumps(key)}: {json.dumps(m, ensure_ascii=False)}{comma}")
    lines += [" }", "}", ""]
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines))


if __name__ == "__main__":
    import argparse

    from puzzle.materials import user_library_path

    parser = argparse.ArgumentParser(description="Build your local Puzzle materials library.")
    parser.add_argument("workbook", type=Path, help="Materials_Hunt_List.xlsx")
    parser.add_argument(
        "--output", type=Path, default=user_library_path(), help="the JSON file to write"
    )
    args = parser.parse_args()
    main(args.workbook, args.output)
